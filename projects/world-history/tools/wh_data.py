"""Shared loading/validation for the World History Explorer data pipeline.

Used by build_data.py (loads -> validates -> writes data/*.json) and
validate_data.py (loads -> validates -> reports, writes nothing). Keeping the
logic here means both scripts see exactly the same rules.

This module is build-time only. Nothing here ships to the browser.

Severity policy (see docs/DATA_MODEL.md #8 and docs/IMPLEMENTATION_PLAN.md
Phase 1):

- ERROR  = structural problem (duplicate/missing ID, broken foreign key,
  malformed year, invalid numeric population, inconsistent low/high bounds).
  The build must fail rather than emit generated JSON that can't be trusted.
- WARNING = content-quality flag (TODO URLs, missing population, low/unknown
  confidence, off-vocabulary values not yet reconciled with docs/DATA_MODEL.md's
  recommended lists, suspicious dates). These are expected at this stage of
  the project and are not blockers.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field

import openpyxl

TODO_URL = "https://en.wikipedia.org/wiki/Time_management#Implementation_of_goals"

HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")

# Recommended/expected controlled vocabularies. Values outside these are
# WARNed on, not rejected -- the workbook is allowed to be ahead of the docs.
CONFIDENCE_LEVELS = {"High", "Medium", "Low", "Unknown"}
POPULATION_STATUS = {"TODO", "Estimated", "Verified"}
POPULATION_BASIS_RECOMMENDED = {
    "polity_total",
    "territorial_segment",
    "modern_country_total",
    "other",
}
# "TBD" is the workbook's own placeholder for basis/method, distinct from the
# recommended vocabulary above -- treated like the TODO URL: expected for now.
PLACEHOLDER_BASIS_METHOD = "TBD"

# docs/DATA_MODEL.md #5 "Recommended relation types". Unlike the vocabularies
# above, docs/IMPLEMENTATION_PLAN.md Phase 1 explicitly lists "invalid
# relation types" among the build-failing checks, so this one is ERROR, not
# WARNING -- see validate_links.
RELATION_TYPES = {
    "successor",
    "partition",
    "unification",
    "conquest",
    "colonization",
    "decolonization",
    "dynastic_transition",
    "institutional_successor",
    "cultural_continuity",
    "migration",
    "other",
}

SHEETS = {
    "Boxes": {"header_row": 1},
    "SEAIs": {"header_row": 1},
    "Population": {"header_row": 3},
    "Regions": {"header_row": 3},
    "Links": {"header_row": 1},
}


@dataclass
class Report:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def error(self, msg: str) -> None:
        self.errors.append(msg)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    @property
    def ok(self) -> bool:
        return not self.errors


def _clean(value):
    """Normalize a raw cell value: strip strings, blank string -> None."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def load_sheet(wb, name: str) -> list[dict]:
    header_row = SHEETS[name]["header_row"]
    ws = wb[name]
    header = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
    rows = []
    for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(c is None for c in raw):
            continue  # skip fully blank rows
        rows.append({col: _clean(v) for col, v in zip(header, raw) if col is not None})
    return rows


def load_workbook(path: str) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    missing = [name for name in SHEETS if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"workbook is missing expected sheet(s): {missing}")
    return {name: load_sheet(wb, name) for name in SHEETS}


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def validate_boxes(boxes: list[dict], report: Report) -> dict[str, dict]:
    by_id: dict[str, dict] = {}
    seen_ids: set[str] = set()
    for row in boxes:
        box_id = row.get("box_id")
        if not box_id:
            report.error(f"Boxes: row missing box_id: {row!r}")
            continue
        if box_id in seen_ids:
            report.error(f"Boxes: duplicate box_id '{box_id}'")
        seen_ids.add(box_id)
        by_id[box_id] = row

        for field_name in ("start_year", "end_year"):
            v = row.get(field_name)
            if v is not None and not _is_number(v):
                report.error(f"Boxes[{box_id}]: malformed {field_name} = {v!r}")

        sy, ey = row.get("start_year"), row.get("end_year")
        if _is_number(sy) and _is_number(ey) and sy > ey:
            report.error(f"Boxes[{box_id}]: start_year ({sy}) > end_year ({ey})")

        color_hex = row.get("color_hex")
        if not color_hex or not HEX_RE.match(str(color_hex)):
            report.error(f"Boxes[{box_id}]: invalid color_hex {color_hex!r}")

        confidence = row.get("confidence")
        if confidence not in CONFIDENCE_LEVELS:
            report.warn(f"Boxes[{box_id}]: confidence {confidence!r} outside {sorted(CONFIDENCE_LEVELS)}")

        url = row.get("url")
        if not url:
            report.error(f"Boxes[{box_id}]: missing url (should at least hold the TODO placeholder)")
        elif url == TODO_URL:
            report.warn(f"Boxes[{box_id}]: url is still the TODO placeholder")

    return by_id


def validate_seais(seais: list[dict], box_ids: set[str], report: Report) -> None:
    seen_ids: set[str] = set()
    for row in seais:
        seai_id = row.get("seai_id")
        if not seai_id:
            report.error(f"SEAIs: row missing seai_id: {row!r}")
            continue
        if seai_id in seen_ids:
            report.error(f"SEAIs: duplicate seai_id '{seai_id}'")
        seen_ids.add(seai_id)

        box_id = row.get("box_id")
        if not box_id:
            report.error(f"SEAIs[{seai_id}]: missing box_id")
        elif box_id not in box_ids:
            report.error(f"SEAIs[{seai_id}]: box_id '{box_id}' does not exist in Boxes")

        for field_name in ("year", "historical_year"):
            v = row.get(field_name)
            if v is not None and not _is_number(v):
                report.error(f"SEAIs[{seai_id}]: malformed {field_name} = {v!r}")

        confidence = row.get("confidence")
        if confidence not in CONFIDENCE_LEVELS:
            report.warn(f"SEAIs[{seai_id}]: confidence {confidence!r} outside {sorted(CONFIDENCE_LEVELS)}")

        url = row.get("url")
        if not url:
            report.error(f"SEAIs[{seai_id}]: missing url (should at least hold the TODO placeholder)")
        elif url == TODO_URL:
            report.warn(f"SEAIs[{seai_id}]: url is still the TODO placeholder")


def validate_population(pop: list[dict], box_ids: set[str], report: Report) -> None:
    seen_ids: set[str] = set()
    seen_knots: dict[tuple, int] = {}
    for row in pop:
        pid = row.get("population_id")
        if not pid:
            report.error(f"Population: row missing population_id: {row!r}")
            continue
        if pid in seen_ids:
            report.error(f"Population: duplicate population_id '{pid}'")
        seen_ids.add(pid)

        box_id = row.get("box_id")
        if not box_id:
            report.error(f"Population[{pid}]: missing box_id")
        elif box_id not in box_ids:
            report.error(f"Population[{pid}]: box_id '{box_id}' does not exist in Boxes")

        year = row.get("year")
        if year is not None and not _is_number(year):
            report.error(f"Population[{pid}]: malformed year = {year!r}")

        knot = (box_id, year, row.get("sample_type"))
        if knot in seen_knots:
            report.warn(f"Population[{pid}]: duplicate knot point {knot} (also row {seen_knots[knot]!r})")
        else:
            seen_knots[knot] = pid

        population = row.get("population")
        low, high = row.get("population_low"), row.get("population_high")
        if population is None:
            report.warn(f"Population[{pid}]: population not yet filled")
        else:
            if not _is_number(population):
                report.error(f"Population[{pid}]: malformed population = {population!r}")
            elif population <= 0:
                report.error(f"Population[{pid}]: population must be > 0, got {population}")
            if _is_number(low) and _is_number(population) and low > population:
                report.error(f"Population[{pid}]: population_low ({low}) > population ({population})")
            if _is_number(high) and _is_number(population) and population > high:
                report.error(f"Population[{pid}]: population ({population}) > population_high ({high})")

        status = row.get("status")
        if status not in POPULATION_STATUS:
            report.warn(f"Population[{pid}]: status {status!r} outside {sorted(POPULATION_STATUS)}")

        for field_name in ("population_basis", "estimate_method"):
            v = row.get(field_name)
            if v is not None and v != PLACEHOLDER_BASIS_METHOD and field_name == "population_basis" and v not in POPULATION_BASIS_RECOMMENDED:
                report.warn(f"Population[{pid}]: {field_name} {v!r} outside recommended vocabulary {sorted(POPULATION_BASIS_RECOMMENDED)}")


def validate_regions(regions: list[dict], report: Report) -> None:
    seen_pairs: set[tuple] = set()
    for row in regions:
        region_id, country = row.get("region_id"), row.get("country")
        if not region_id or not row.get("region_name") or not country:
            report.error(f"Regions: row missing region_id/region_name/country: {row!r}")
            continue
        pair = (region_id, country)
        if pair in seen_pairs:
            report.error(f"Regions: duplicate (region_id, country) pair {pair}")
        seen_pairs.add(pair)


def validate_links(links: list[dict], boxes_by_id: dict[str, dict], report: Report) -> None:
    seen_ids: set[str] = set()
    for row in links:
        link_id = row.get("link_id")
        if not link_id:
            report.error(f"Links: row missing link_id: {row!r}")
            continue
        if link_id in seen_ids:
            report.error(f"Links: duplicate link_id '{link_id}'")
        seen_ids.add(link_id)

        source_id, target_id = row.get("source_box_id"), row.get("target_box_id")
        source_box = boxes_by_id.get(source_id) if source_id else None
        target_box = boxes_by_id.get(target_id) if target_id else None
        if not source_id:
            report.error(f"Links[{link_id}]: missing source_box_id")
        elif source_box is None:
            report.error(f"Links[{link_id}]: source_box_id '{source_id}' does not exist in Boxes")
        if not target_id:
            report.error(f"Links[{link_id}]: missing target_box_id")
        elif target_box is None:
            report.error(f"Links[{link_id}]: target_box_id '{target_id}' does not exist in Boxes")

        year = row.get("year")
        if year is not None and not _is_number(year):
            report.error(f"Links[{link_id}]: malformed year = {year!r}")

        # docs/IMPLEMENTATION_PLAN.md Phase 1 explicitly fails the build on
        # "invalid relation types" -- ERROR, unlike the WARN-level vocabulary
        # checks elsewhere in this file (see RELATION_TYPES comment above).
        relation_type = row.get("relation_type")
        if relation_type not in RELATION_TYPES:
            report.error(f"Links[{link_id}]: relation_type {relation_type!r} not in {sorted(RELATION_TYPES)}")

        # "Transition year sensible" (docs/DATA_MODEL.md #8) is a content-
        # quality check, not structural -- WARN rather than ERROR.
        if _is_number(year) and source_box is not None and target_box is not None:
            lo = min(v for v in (source_box.get("start_year"), target_box.get("start_year")) if _is_number(v))
            hi = max(v for v in (source_box.get("end_year"), target_box.get("end_year")) if _is_number(v))
            if not (lo <= year <= hi):
                report.warn(
                    f"Links[{link_id}]: year {year} falls outside the source/target boxes' "
                    f"combined span [{lo}, {hi}]"
                )

        confidence = row.get("confidence")
        if confidence not in CONFIDENCE_LEVELS:
            report.warn(f"Links[{link_id}]: confidence {confidence!r} outside {sorted(CONFIDENCE_LEVELS)}")

        url = row.get("source_url")
        if not url:
            report.warn(f"Links[{link_id}]: missing source_url")
        elif url == TODO_URL:
            report.warn(f"Links[{link_id}]: source_url is still the TODO placeholder")


def compute_log10_population(pop: list[dict], report: Report) -> None:
    """Recompute log10_population from population at build time -- it's a
    derived field per docs/DATA_MODEL.md, not hand-authored."""
    for row in pop:
        population = row.get("population")
        if _is_number(population) and population > 0:
            computed = round(math.log10(population), 4)
            existing = row.get("log10_population")
            if _is_number(existing) and abs(existing - computed) > 1e-3:
                report.warn(
                    f"Population[{row.get('population_id')}]: workbook log10_population "
                    f"{existing} does not match computed {computed}; overwriting with computed value"
                )
            row["log10_population"] = computed
        else:
            row["log10_population"] = None


def validate_all(tables: dict[str, list[dict]]) -> Report:
    report = Report()
    boxes_by_id = validate_boxes(tables["Boxes"], report)
    box_ids = set(boxes_by_id)
    validate_seais(tables["SEAIs"], box_ids, report)
    validate_population(tables["Population"], box_ids, report)
    validate_regions(tables["Regions"], report)
    validate_links(tables["Links"], boxes_by_id, report)
    compute_log10_population(tables["Population"], report)
    return report
